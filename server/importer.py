"""قراءة ملفي Excel (الطلاب + المواد) وتحويلهما إلى قوائم جاهزة للإدخال في قاعدة البيانات."""
import re
import openpyxl

STAGES = {"الأولى": 1, "الثانية": 2, "الثالثة": 3, "الرابعة": 4, "الخامسة": 5, "السادسة": 6}


def _clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def _stage(text):
    for word, num in STAGES.items():
        if word in text:
            return num
    return None


def _shift(text):
    if "صباح" in text:
        return "morning"
    if "مسا" in text:
        return "evening"
    return None


def _header_map(ws, needle):
    """يبحث عن صف العناوين (الذي يحتوي needle) ويعيد (رقم الصف، {اسم العمود: فهرسه})."""
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        cells = [_clean(c) for c in row]
        if needle in cells:
            return r_idx, {name: i for i, name in enumerate(cells) if name}
    return None, None


def parse_students(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = [wb["كل الطلاب"]] if "كل الطلاب" in wb.sheetnames else [wb[n] for n in wb.sheetnames]
    found = {}
    for ws in sheets:
        hdr_row, cols = _header_map(ws, "الرقم المعرفي")
        if not cols:
            continue
        for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
            uid = _clean(row[cols["الرقم المعرفي"]]) if cols["الرقم المعرفي"] < len(row) else ""
            name = _clean(row[cols["اسم الطالب"]]) if "اسم الطالب" in cols else ""
            if not uid or not name:
                continue
            stage = _stage(_clean(row[cols["المرحلة"]])) if "المرحلة" in cols else _stage(ws.title)
            shift = _shift(_clean(row[cols["الدوام"]])) if "الدوام" in cols else _shift(ws.title)
            prev = found.get(uid)
            # الطالب الموجود في أكثر من مرحلة يُعتمد في المرحلة الأعلى
            if prev is None or (stage or 0) > (prev["stage"] or 0):
                found[uid] = {"uid": uid, "name": name, "stage": stage, "shift": shift}
    return list(found.values())


def parse_courses(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["جميع المواد"] if "جميع المواد" in wb.sheetnames else wb[wb.sheetnames[0]]
    hdr_row, cols = _header_map(ws, "رمز المقرر")
    if not cols:
        raise ValueError("لم أجد صف العناوين في ملف المواد")
    out = []
    for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
        def g(name):
            i = cols.get(name)
            return _clean(row[i]) if i is not None and i < len(row) else ""
        code, ref = g("رمز المقرر"), g("معرف المادة")
        if not code or not ref:
            continue
        season_txt = g("فصل السنة")
        out.append({
            "ref_id": ref,
            "dept": ref.split("|")[0],
            "dept_name": g("القسم"),
            "stage": _stage(g("المرحلة")),
            "semester": int(g("الفصل الدراسي") or 0),
            "season": 1 if "الأول" in season_txt else 2,
            "code": code,
            "name_ar": g("اسم المادة بالعربية"),
            "name_en": g("اسم المادة بالإنجليزية"),
            "ects": int(float(g("ECTS") or 0)),
        })
    return out
