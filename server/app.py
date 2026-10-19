"""خادم نظام الحضور والغياب — جامعة المعقل الأهلية (البصرة)

Flask + SQLite. يوفّر واجهة JSON للتطبيق، ويصدّر كشوف Excel.
"""
import hashlib
import base64
import hmac
import html
import io
import json
import os
import re
import secrets
import sqlite3
import time
import urllib.parse
import urllib.request
from functools import wraps

from flask import Flask, g, jsonify, request, send_file
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from werkzeug.security import check_password_hash, generate_password_hash

BASE = os.path.dirname(os.path.abspath(__file__))
MAX_SESSION_SECONDS = 4 * 3600          # تُغلق المحاضرة تلقائياً بعد 4 ساعات
TOKEN_MAX_AGE = 30 * 24 * 3600          # صلاحية الدخول 30 يوماً
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
UID_RE = re.compile(r"^\d{10}$")
DEFAULT_SETTINGS = {
    "current_season": "1",      # الفصل الحالي: 1 = الأول، 2 = الثاني
    "late_minutes": "15",       # بعد كم دقيقة يُسجَّل الطالب «متأخراً»
    "absence_limit_pct": "20",  # نسبة الغياب التي يُنذر الطالب عندها
    "qr_step": "20",            # مدة بقاء رمز QR قبل أن يتبدل (ثوانٍ)
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS users(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  role TEXT NOT NULL CHECK(role IN ('admin','doctor','student')),
  uid TEXT UNIQUE,
  email TEXT UNIQUE,
  name TEXT NOT NULL,
  pw_hash TEXT,
  activated INTEGER NOT NULL DEFAULT 0,
  must_change INTEGER NOT NULL DEFAULT 0,
  device_id TEXT,
  department TEXT,
  stage INTEGER,
  shift TEXT,
  active INTEGER NOT NULL DEFAULT 1,
  token_ver INTEGER NOT NULL DEFAULT 0,
  created_at INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS courses(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  ref_id TEXT UNIQUE NOT NULL,
  dept TEXT NOT NULL,
  dept_name TEXT NOT NULL,
  stage INTEGER NOT NULL,
  semester INTEGER NOT NULL,
  season INTEGER NOT NULL,
  code TEXT NOT NULL,
  name_ar TEXT NOT NULL,
  name_en TEXT,
  ects INTEGER
);
CREATE TABLE IF NOT EXISTS assignments(
  doctor_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  course_id INTEGER NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
  PRIMARY KEY(doctor_id, course_id)
);
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS oauth_pending(nonce TEXT PRIMARY KEY, device TEXT, token TEXT, error TEXT, created INTEGER NOT NULL);
CREATE TABLE IF NOT EXISTS sessions(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  course_id INTEGER NOT NULL REFERENCES courses(id) ON DELETE CASCADE,
  doctor_id INTEGER NOT NULL REFERENCES users(id),
  shift TEXT NOT NULL DEFAULT 'all',
  secret TEXT NOT NULL,
  late_after INTEGER NOT NULL,
  started_at INTEGER NOT NULL,
  ended_at INTEGER
);
CREATE INDEX IF NOT EXISTS ix_sessions_course ON sessions(course_id);
CREATE TABLE IF NOT EXISTS attendance(
  session_id INTEGER NOT NULL REFERENCES sessions(id) ON DELETE CASCADE,
  student_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  status TEXT NOT NULL CHECK(status IN ('present','late','excused')),
  at INTEGER NOT NULL,
  method TEXT NOT NULL DEFAULT 'qr',
  PRIMARY KEY(session_id, student_id)
);
CREATE INDEX IF NOT EXISTS ix_students_group ON users(role, department, stage);
"""

_FAILS = {}  # تحديد محاولات الدخول الخاطئة: المفتاح -> (عدد، وقت أول محاولة)


# ---------------------------------------------------------------- أدوات عامة
def now():
    return int(time.time())


def err(code, message, status=400):
    return jsonify({"error": code, "message": message}), status


def _too_many(key, limit, window):
    n, first = _FAILS.get(key, (0, 0))
    if n and now() - first > window:
        _FAILS.pop(key, None)
        return False
    return n >= limit


def _fail(key):
    n, first = _FAILS.get(key, (0, now()))
    _FAILS[key] = (n + 1, first)


def _clear(key):
    _FAILS.pop(key, None)


def norm_email(v):
    v = (v or "").strip().lower()
    return v if EMAIL_RE.match(v) else None


def google_exchange(code, redirect_uri):
    """يبدّل رمز Google بالبريد الموثّق. يُستبدل في الاختبارات."""
    data = urllib.parse.urlencode({
        "code": code, "client_id": os.environ["MAQAL_GOOGLE_CLIENT_ID"],
        "client_secret": os.environ["MAQAL_GOOGLE_CLIENT_SECRET"],
        "redirect_uri": redirect_uri, "grant_type": "authorization_code"}).encode()
    with urllib.request.urlopen(urllib.request.Request("https://oauth2.googleapis.com/token", data=data), timeout=15) as r:
        tok = json.loads(r.read().decode())
    payload = tok["id_token"].split(".")[1]
    info = json.loads(base64.urlsafe_b64decode(payload + "=" * (-len(payload) % 4)))
    if not info.get("email_verified"):
        return None
    return (info.get("email") or "").lower()


def create_app(db_path=None, secret=None):
    app = Flask(__name__)
    app.config["DB_PATH"] = db_path or os.environ.get("MAQAL_DB") or os.path.join(BASE, "instance", "attendance.db")
    os.makedirs(os.path.dirname(app.config["DB_PATH"]), exist_ok=True)
    app.config["SECRET"] = secret or _load_secret(os.path.dirname(app.config["DB_PATH"]))
    app.json.ensure_ascii = False
    app.json.sort_keys = False
    init_db(app.config["DB_PATH"])

    auth_s = URLSafeTimedSerializer(app.config["SECRET"], salt="auth")
    export_s = URLSafeTimedSerializer(app.config["SECRET"], salt="export")

    # ---------------------------------------------------------- قاعدة البيانات
    def db():
        if "db" not in g:
            conn = sqlite3.connect(app.config["DB_PATH"], timeout=15)
            conn.row_factory = sqlite3.Row
            conn.execute("PRAGMA foreign_keys=ON")
            g.db = conn
        return g.db

    @app.teardown_appcontext
    def _close(_exc):
        conn = g.pop("db", None)
        if conn is not None:
            conn.close()

    def setting(key):
        row = db().execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
        return int(row["value"]) if row else int(DEFAULT_SETTINGS[key])

    def autoclose():
        db().execute(
            "UPDATE sessions SET ended_at=started_at+? WHERE ended_at IS NULL AND started_at<?",
            (MAX_SESSION_SECONDS, now() - MAX_SESSION_SECONDS),
        )
        db().commit()

    # ---------------------------------------------------------- المصادقة
    def issue_token(user):
        return auth_s.dumps({"u": user["id"], "v": user["token_ver"]})

    def public_user(u):
        return {
            "id": u["id"], "role": u["role"], "name": u["name"], "uid": u["uid"], "email": u["email"],
            "department": u["department"], "stage": u["stage"], "shift": u["shift"],
            "must_change": bool(u["must_change"]),
        }

    def auth(*roles):
        def deco(fn):
            @wraps(fn)
            def wrapper(*a, **kw):
                head = request.headers.get("Authorization", "")
                if not head.startswith("Bearer "):
                    return err("unauthorized", "يلزم تسجيل الدخول", 401)
                try:
                    data = auth_s.loads(head[7:], max_age=TOKEN_MAX_AGE)
                except (BadSignature, SignatureExpired):
                    return err("unauthorized", "انتهت الجلسة، سجّل الدخول من جديد", 401)
                u = db().execute("SELECT * FROM users WHERE id=?", (data["u"],)).fetchone()
                if not u or not u["active"] or u["token_ver"] != data["v"]:
                    return err("unauthorized", "انتهت الجلسة، سجّل الدخول من جديد", 401)
                if roles and u["role"] not in roles:
                    return err("forbidden", "لا تملك صلاحية لهذا الإجراء", 403)
                g.user = u
                return fn(*a, **kw)
            return wrapper
        return deco

    def body():
        return request.get_json(silent=True) or {}

    # ---------------------------------------------------------- QR الدوّار
    def code_for(secret_, sid, slot):
        digest = hmac.new(secret_.encode(), f"{sid}:{slot}".encode(), hashlib.sha256).digest()
        return f"{int.from_bytes(digest[:4], 'big') % 1_000_000:06d}"

    def current_qr(sess):
        step = setting("qr_step")
        slot = int(time.time() // step)
        code = code_for(sess["secret"], sess["id"], slot)
        return {"payload": f"MQL|{sess['id']}|{code}", "code": code,
                "expires_in": step - int(time.time() % step), "step": step}

    def code_valid(sess, code):
        step = setting("qr_step")
        slot = int(time.time() // step)
        return any(hmac.compare_digest(code_for(sess["secret"], sess["id"], s), code) for s in (slot, slot - 1))

    # ---------------------------------------------------------- منطق الحضور
    def enrolled(course, shift="all"):
        sql = "SELECT * FROM users WHERE role='student' AND active=1 AND department=? AND stage=?"
        args = [course["dept"], course["stage"]]
        if shift in ("morning", "evening"):
            sql += " AND shift=?"
            args.append(shift)
        return db().execute(sql + " ORDER BY name", args).fetchall()

    def unassigned(course, shift="all"):
        """طلاب مرحلة المادة الذين لم يُحدَّد قسمهم بعد (لن يظهروا في الكشف حتى يفعّلوا حساباتهم)."""
        sql = "SELECT COUNT(*) FROM users WHERE role='student' AND active=1 AND stage=? AND (department IS NULL OR department='')"
        args = [course["stage"]]
        if shift in ("morning", "evening"):
            sql += " AND shift=?"
            args.append(shift)
        return db().execute(sql, args).fetchone()[0]

    def course_dict(c):
        return {"id": c["id"], "code": c["code"], "name": c["name_ar"], "name_en": c["name_en"],
                "dept": c["dept"], "dept_name": c["dept_name"], "stage": c["stage"],
                "semester": c["semester"], "season": c["season"], "ects": c["ects"]}

    def summarize(student, sessions, records):
        """sessions: الجلسات المنتهية للمادة. records: {session_id: row} لهذا الطالب."""
        applicable = [s for s in sessions if s["shift"] in ("all", student["shift"])]
        p = sum(1 for s in applicable if s["id"] in records and records[s["id"]]["status"] == "present")
        l = sum(1 for s in applicable if s["id"] in records and records[s["id"]]["status"] == "late")
        e = sum(1 for s in applicable if s["id"] in records and records[s["id"]]["status"] == "excused")
        total = len(applicable)
        absent = total - p - l - e
        denom = total - e
        pct = None if denom <= 0 else round((p + l) * 100 / denom)
        return {"total": total, "present": p, "late": l, "excused": e, "absent": absent, "pct": pct}

    def course_report(course):
        sessions = db().execute(
            "SELECT * FROM sessions WHERE course_id=? AND ended_at IS NOT NULL ORDER BY started_at", (course["id"],)
        ).fetchall()
        students = enrolled(course)
        recs = {}
        if sessions:
            q = ",".join("?" * len(sessions))
            for r in db().execute(f"SELECT * FROM attendance WHERE session_id IN ({q})", [s["id"] for s in sessions]):
                recs.setdefault(r["student_id"], {})[r["session_id"]] = r
        limit = setting("absence_limit_pct")
        rows = []
        for st in students:
            sm = summarize(st, sessions, recs.get(st["id"], {}))
            sm.update({"student_id": st["id"], "uid": st["uid"], "name": st["name"], "shift": st["shift"],
                       "warn": sm["pct"] is not None and (100 - sm["pct"]) >= limit})
            rows.append(sm)
        return sessions, rows, recs

    def can_access_course(user, course_id):
        if user["role"] == "admin":
            return True
        return db().execute("SELECT 1 FROM assignments WHERE doctor_id=? AND course_id=?",
                            (user["id"], course_id)).fetchone() is not None

    # =================================================================== عام
    @app.get("/")
    def index():
        return jsonify({"app": "maqal-attendance", "ok": True})

    @app.get("/api/ping")
    def ping():
        return jsonify({"ok": True, "app": "maqal-attendance", "version": 1})

    @app.get("/api/meta/departments")
    def departments():
        rows = db().execute("SELECT DISTINCT dept, dept_name FROM courses ORDER BY dept_name").fetchall()
        return jsonify([{"key": r["dept"], "name": r["dept_name"]} for r in rows])

    # =================================================================== الدخول
    @app.post("/api/auth/login")
    def login():
        d = body()
        ident = (d.get("identifier") or "").strip().lower()
        pw = d.get("password") or ""
        device = (d.get("device_id") or "").strip()
        key = f"login:{request.remote_addr}:{ident}"
        if _too_many(key, 6, 300):
            return err("too_many", "محاولات كثيرة خاطئة، حاول بعد 5 دقائق", 429)
        u = None
        if UID_RE.match(ident):
            u = db().execute("SELECT * FROM users WHERE uid=?", (ident,)).fetchone()
        elif "@" in ident:
            u = db().execute("SELECT * FROM users WHERE email=?", (ident,)).fetchone()
        if u and u["role"] == "student" and not u["activated"]:
            return err("not_activated", "هذا الحساب لم يُفعَّل بعد. اختر «تفعيل حساب طالب».", 409)
        if not u or not u["pw_hash"] or not check_password_hash(u["pw_hash"], pw):
            _fail(key)
            return err("bad_credentials", "بيانات الدخول غير صحيحة", 401)
        if not u["active"]:
            return err("disabled", "الحساب موقوف، راجع إدارة الجامعة", 403)
        if u["role"] == "student":
            if not device:
                return err("device_required", "تعذّر تحديد الجهاز", 400)
            if u["device_id"] and u["device_id"] != device:
                return err("device_mismatch", "هذا الحساب مرتبط بهاتف آخر. راجع إدارة الجامعة لفك الارتباط.", 403)
            if not u["device_id"]:
                db().execute("UPDATE users SET device_id=? WHERE id=?", (device, u["id"]))
                db().commit()
        _clear(key)
        want = (d.get("role") or "").strip()
        if want and want != u["role"]:
            names = {"student": "طالب", "doctor": "دكتور", "admin": "مدير"}
            return err("wrong_role", f"هذا الحساب ليس حساب {names.get(want, want)}. اختر الدور الصحيح.", 403)
        return jsonify({"token": issue_token(u), "user": public_user(u)})

    # ---- الدخول بحساب Google (عبر المتصفح، والتطبيق يستعلم بالـ nonce)
    def public_base():
        return (os.environ.get("MAQAL_PUBLIC_URL") or request.url_root).rstrip("/")

    @app.get("/api/auth/options")
    def auth_options():
        return jsonify({"google": bool(os.environ.get("MAQAL_GOOGLE_CLIENT_ID"))})

    def done_page(ok, msg):
        color = "#16a34a" if ok else "#dc2626"
        return (f'<!doctype html><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
                f'<body style="font-family:sans-serif;text-align:center;padding:48px 20px;direction:rtl">'
                f'<h2 style="color:{color}">{"تم" if ok else "تعذّر الدخول"}</h2><p>{html.escape(msg)}</p></body>')

    @app.get("/auth/google/start")
    def google_start():
        cid = os.environ.get("MAQAL_GOOGLE_CLIENT_ID")
        nonce = (request.args.get("nonce") or "").strip()
        if not cid:
            return done_page(False, "الدخول بحساب Google غير مفعّل على هذا الخادم."), 503
        if not 16 <= len(nonce) <= 64:
            return done_page(False, "طلب غير صالح."), 400
        db().execute("DELETE FROM oauth_pending WHERE created<?", (now() - 600,))
        db().execute("INSERT OR REPLACE INTO oauth_pending(nonce,device,created) VALUES(?,?,?)",
                     (nonce, (request.args.get("device") or "").strip(), now()))
        db().commit()
        q = urllib.parse.urlencode({
            "client_id": cid, "redirect_uri": public_base() + "/auth/google/callback", "response_type": "code",
            "scope": "openid email", "state": nonce, "prompt": "select_account"})
        return "", 302, {"Location": "https://accounts.google.com/o/oauth2/v2/auth?" + q}

    @app.get("/auth/google/callback")
    def google_callback():
        nonce = request.args.get("state") or ""
        row = db().execute("SELECT * FROM oauth_pending WHERE nonce=?", (nonce,)).fetchone()
        if not row or request.args.get("error") or not request.args.get("code"):
            return done_page(False, "أُلغي الدخول أو انتهت صلاحية الطلب. ارجع للتطبيق وأعد المحاولة."), 400

        def fail(m):
            db().execute("UPDATE oauth_pending SET error=? WHERE nonce=?", (m, nonce))
            db().commit()
            return done_page(False, m), 403
        try:
            email = google_exchange(request.args["code"], public_base() + "/auth/google/callback")
        except Exception:
            return fail("تعذّر التحقق من حساب Google.")
        if not email:
            return fail("بريد Google غير موثّق.")
        u = db().execute("SELECT * FROM users WHERE email=?", (email,)).fetchone()
        if not u:
            return fail(f"البريد {email} غير مرتبط بأي حساب. اطلب من الإدارة ربطه، أو فعّل حسابك واكتب هذا البريد.")
        if u["role"] == "student" and not u["activated"]:
            return fail("هذا الحساب لم يُفعَّل بعد.")
        if not u["active"]:
            return fail("الحساب موقوف، راجع إدارة الجامعة.")
        if u["role"] == "student":
            device = row["device"]
            if not device:
                return fail("تعذّر تحديد الجهاز.")
            if u["device_id"] and u["device_id"] != device:
                return fail("هذا الحساب مرتبط بهاتف آخر. راجع إدارة الجامعة.")
            if not u["device_id"]:
                db().execute("UPDATE users SET device_id=? WHERE id=?", (device, u["id"]))
        db().execute("UPDATE oauth_pending SET token=? WHERE nonce=?", (issue_token(u), nonce))
        db().commit()
        return done_page(True, "ارجع الآن إلى تطبيق حضور المعقل.")

    @app.get("/api/auth/google/poll")
    def google_poll():
        nonce = request.args.get("nonce") or ""
        row = db().execute("SELECT * FROM oauth_pending WHERE nonce=?", (nonce,)).fetchone()
        if not row:
            return err("expired", "انتهت صلاحية الطلب", 404)
        if row["error"]:
            db().execute("DELETE FROM oauth_pending WHERE nonce=?", (nonce,))
            db().commit()
            return err("google_failed", row["error"], 403)
        if not row["token"]:
            return jsonify({"status": "pending"})
        want = (request.args.get("role") or "").strip()
        data = auth_s.loads(row["token"])
        u = db().execute("SELECT * FROM users WHERE id=?", (data["u"],)).fetchone()
        db().execute("DELETE FROM oauth_pending WHERE nonce=?", (nonce,))
        db().commit()
        if want and want != u["role"]:
            return err("wrong_role", "هذا الحساب ليس من الدور المختار. اختر الدور الصحيح.", 403)
        return jsonify({"status": "ok", "token": row["token"], "user": public_user(u)})

    @app.post("/api/auth/activate")
    def activate():
        d = body()
        uid = (d.get("uid") or "").strip()
        pw = d.get("password") or ""
        dept = (d.get("department") or "").strip()
        email = norm_email(d.get("email")) if d.get("email") else None
        device = (d.get("device_id") or "").strip()
        key = f"act:{request.remote_addr}"
        if _too_many(key, 15, 600):
            return err("too_many", "محاولات كثيرة، حاول لاحقاً", 429)
        if not UID_RE.match(uid):
            _fail(key)
            return err("bad_uid", "الرقم المعرفي يجب أن يكون 10 أرقام")
        u = db().execute("SELECT * FROM users WHERE uid=? AND role='student'", (uid,)).fetchone()
        if not u:
            _fail(key)
            return err("unknown_uid", "الرقم المعرفي غير موجود في قاعدة الطلاب", 404)
        if u["activated"]:
            return err("already_active", "الحساب مفعّل مسبقاً، سجّل الدخول. وإن لم تكن أنت من فعّله فراجع الإدارة.", 409)
        if len(pw) < 6:
            return err("weak_password", "كلمة المرور يجب ألا تقل عن 6 أحرف")
        if d.get("email") and not email:
            return err("bad_email", "صيغة البريد الإلكتروني غير صحيحة")
        if not device:
            return err("device_required", "تعذّر تحديد الجهاز")
        if not db().execute("SELECT 1 FROM courses WHERE dept=?", (dept,)).fetchone():
            return err("bad_department", "اختر القسم")
        if email and db().execute("SELECT 1 FROM users WHERE email=?", (email,)).fetchone():
            return err("email_taken", "هذا البريد مستخدم لحساب آخر", 409)
        db().execute(
            "UPDATE users SET pw_hash=?, activated=1, email=?, department=?, device_id=? WHERE id=?",
            (generate_password_hash(pw), email, dept, device, u["id"]),
        )
        db().commit()
        u = db().execute("SELECT * FROM users WHERE id=?", (u["id"],)).fetchone()
        return jsonify({"token": issue_token(u), "user": public_user(u)})

    @app.get("/api/me")
    @auth()
    def me():
        return jsonify(public_user(g.user))

    @app.post("/api/auth/change-password")
    @auth()
    def change_password():
        d = body()
        old, new = d.get("old") or "", d.get("new") or ""
        if len(new) < 6:
            return err("weak_password", "كلمة المرور الجديدة يجب ألا تقل عن 6 أحرف")
        if not g.user["pw_hash"] or not check_password_hash(g.user["pw_hash"], old):
            return err("bad_credentials", "كلمة المرور الحالية غير صحيحة", 401)
        db().execute("UPDATE users SET pw_hash=?, must_change=0 WHERE id=?",
                     (generate_password_hash(new), g.user["id"]))
        db().commit()
        return jsonify({"ok": True})

    @app.post("/api/auth/set-email")
    @auth()
    def set_email():
        email = norm_email(body().get("email"))
        if not email:
            return err("bad_email", "صيغة البريد الإلكتروني غير صحيحة")
        other = db().execute("SELECT id FROM users WHERE email=?", (email,)).fetchone()
        if other and other["id"] != g.user["id"]:
            return err("email_taken", "هذا البريد مستخدم لحساب آخر", 409)
        db().execute("UPDATE users SET email=? WHERE id=?", (email, g.user["id"]))
        db().commit()
        return jsonify({"ok": True, "email": email})

    # =================================================================== الطالب
    def student_courses(u):
        if not u["department"]:
            return []
        return db().execute(
            "SELECT * FROM courses WHERE dept=? AND stage=? AND season=? ORDER BY semester, name_ar",
            (u["department"], u["stage"], setting("current_season")),
        ).fetchall()

    def student_records(u, course_id):
        sessions = db().execute(
            "SELECT * FROM sessions WHERE course_id=? AND ended_at IS NOT NULL ORDER BY started_at", (course_id,)
        ).fetchall()
        recs = {r["session_id"]: r for r in db().execute(
            "SELECT a.* FROM attendance a JOIN sessions s ON s.id=a.session_id WHERE s.course_id=? AND a.student_id=?",
            (course_id, u["id"]))}
        return sessions, recs

    @app.get("/api/student/home")
    @auth("student")
    def student_home():
        autoclose()
        u = g.user
        limit = setting("absence_limit_pct")
        out, tp, ta = [], 0, 0
        for c in student_courses(u):
            sessions, recs = student_records(u, c["id"])
            sm = summarize(u, sessions, recs)
            active = db().execute(
                "SELECT id FROM sessions WHERE course_id=? AND ended_at IS NULL AND shift IN ('all',?)",
                (c["id"], u["shift"])).fetchone()
            sm["warn"] = sm["pct"] is not None and (100 - sm["pct"]) >= limit
            sm["live"] = active is not None
            out.append({**course_dict(c), **sm})
            tp += sm["present"] + sm["late"]
            ta += sm["absent"]
        total = sum(c["total"] for c in out)
        overall = None if not total else round(tp * 100 / max(1, total - sum(c["excused"] for c in out)))
        return jsonify({"user": public_user(u), "courses": out, "overall_pct": overall,
                        "absences": ta, "limit_pct": limit,
                        "needs_department": not u["department"]})

    @app.get("/api/student/courses/<int:cid>/history")
    @auth("student")
    def student_history(cid):
        u = g.user
        c = db().execute("SELECT * FROM courses WHERE id=?", (cid,)).fetchone()
        if not c or c["dept"] != u["department"] or c["stage"] != u["stage"]:
            return err("not_found", "المادة غير موجودة", 404)
        sessions, recs = student_records(u, cid)
        rows = []
        for s in reversed([s for s in sessions if s["shift"] in ("all", u["shift"])]):
            r = recs.get(s["id"])
            rows.append({"session_id": s["id"], "date": s["started_at"],
                         "status": r["status"] if r else "absent", "at": r["at"] if r else None})
        return jsonify({"course": course_dict(c), "summary": summarize(u, sessions, recs), "rows": rows})

    @app.post("/api/attendance/checkin")
    @auth("student")
    def checkin():
        autoclose()
        u, d = g.user, body()
        code = re.sub(r"\D", "", str(d.get("code") or ""))
        device = (d.get("device_id") or "").strip()
        key = f"ci:{u['id']}"
        if _too_many(key, 8, 60):
            return err("too_many", "محاولات خاطئة كثيرة، انتظر دقيقة ثم أعد المحاولة", 429)
        if u["device_id"] and device != u["device_id"]:
            return err("device_mismatch", "هذا الحساب مرتبط بهاتف آخر", 403)
        if len(code) != 6:
            _fail(key)
            return err("bad_code", "الرمز يتكون من 6 أرقام")
        sid = d.get("session_id")
        q = ("SELECT s.*, c.name_ar, c.dept, c.stage FROM sessions s JOIN courses c ON c.id=s.course_id "
             "WHERE s.ended_at IS NULL AND c.dept=? AND c.stage=? AND s.shift IN ('all',?)")
        args = [u["department"], u["stage"], u["shift"]]
        if sid:
            q += " AND s.id=?"
            args.append(int(sid))
        match = next((s for s in db().execute(q, args).fetchall() if code_valid(s, code)), None)
        if not match:
            _fail(key)
            return err("invalid_code", "الرمز غير صحيح أو انتهت صلاحيته. امسح الرمز الحالي من جديد.")
        _clear(key)
        prev = db().execute("SELECT status FROM attendance WHERE session_id=? AND student_id=?",
                            (match["id"], u["id"])).fetchone()
        if prev:
            return jsonify({"ok": True, "already": True, "status": prev["status"], "course": match["name_ar"]})
        status = "late" if now() - match["started_at"] > match["late_after"] * 60 else "present"
        db().execute("INSERT INTO attendance(session_id,student_id,status,at,method) VALUES(?,?,?,?,'qr')",
                     (match["id"], u["id"], status, now()))
        db().commit()
        return jsonify({"ok": True, "already": False, "status": status, "course": match["name_ar"]})

    # =================================================================== الدكتور
    @app.get("/api/doctor/courses")
    @auth("doctor")
    def doctor_courses():
        autoclose()
        rows = db().execute(
            "SELECT c.* FROM courses c JOIN assignments a ON a.course_id=c.id WHERE a.doctor_id=? "
            "ORDER BY c.stage, c.name_ar", (g.user["id"],)).fetchall()
        out = []
        for c in rows:
            n_students = len(enrolled(c))
            n_sessions = db().execute("SELECT COUNT(*) n FROM sessions WHERE course_id=? AND ended_at IS NOT NULL",
                                      (c["id"],)).fetchone()["n"]
            act = db().execute("SELECT id FROM sessions WHERE course_id=? AND ended_at IS NULL", (c["id"],)).fetchone()
            out.append({**course_dict(c), "students": n_students, "sessions": n_sessions,
                        "active_session": act["id"] if act else None})
        return jsonify(out)

    @app.post("/api/doctor/sessions")
    @auth("doctor")
    def start_session():
        autoclose()
        d = body()
        cid = d.get("course_id")
        shift = d.get("shift", "all")
        if shift not in ("all", "morning", "evening"):
            return err("bad_shift", "نوع الدوام غير صحيح")
        c = db().execute("SELECT * FROM courses WHERE id=?", (cid,)).fetchone()
        if not c or not can_access_course(g.user, c["id"]):
            return err("forbidden", "هذه المادة غير مسندة إليك", 403)
        act = db().execute("SELECT * FROM sessions WHERE course_id=? AND ended_at IS NULL", (c["id"],)).fetchone()
        if act:
            return jsonify({"session_id": act["id"], "resumed": True})
        late = d.get("late_after")
        late = int(late) if str(late or "").isdigit() else setting("late_minutes")
        cur = db().execute(
            "INSERT INTO sessions(course_id,doctor_id,shift,secret,late_after,started_at) VALUES(?,?,?,?,?,?)",
            (c["id"], g.user["id"], shift, secrets.token_hex(16), late, now()))
        db().commit()
        return jsonify({"session_id": cur.lastrowid, "resumed": False})

    def load_session(sid):
        s = db().execute("SELECT * FROM sessions WHERE id=?", (sid,)).fetchone()
        if not s or not can_access_course(g.user, s["course_id"]):
            return None
        return s

    def roster(s):
        c = db().execute("SELECT * FROM courses WHERE id=?", (s["course_id"],)).fetchone()
        recs = {r["student_id"]: r for r in db().execute("SELECT * FROM attendance WHERE session_id=?", (s["id"],))}
        rows = []
        for st in enrolled(c, s["shift"]):
            r = recs.get(st["id"])
            rows.append({"student_id": st["id"], "uid": st["uid"], "name": st["name"], "shift": st["shift"],
                         "status": r["status"] if r else "absent", "at": r["at"] if r else None,
                         "method": r["method"] if r else None})
        return c, rows

    @app.get("/api/doctor/sessions/<int:sid>")
    @auth("doctor", "admin")
    def session_state(sid):
        autoclose()
        s = load_session(sid)
        if not s:
            return err("not_found", "المحاضرة غير موجودة", 404)
        c, rows = roster(s)
        out = {"id": s["id"], "course": course_dict(c), "shift": s["shift"], "started_at": s["started_at"],
               "ended_at": s["ended_at"], "late_after": s["late_after"], "server_time": now(),
               "counts": {"total": len(rows),
                          "present": sum(1 for r in rows if r["status"] == "present"),
                          "late": sum(1 for r in rows if r["status"] == "late"),
                          "excused": sum(1 for r in rows if r["status"] == "excused")},
               "unassigned": unassigned(c, s["shift"]),
               "roster": rows}
        if s["ended_at"] is None:
            out["qr"] = current_qr(s)
        return jsonify(out)

    @app.post("/api/doctor/sessions/<int:sid>/mark")
    @auth("doctor", "admin")
    def mark(sid):
        s = load_session(sid)
        if not s:
            return err("not_found", "المحاضرة غير موجودة", 404)
        d = body()
        status = d.get("status")
        if status not in ("present", "late", "excused", "absent"):
            return err("bad_status", "حالة غير صحيحة")
        c = db().execute("SELECT * FROM courses WHERE id=?", (s["course_id"],)).fetchone()
        if not any(r["id"] == d.get("student_id") for r in enrolled(c, s["shift"])):
            return err("not_enrolled", "الطالب غير مسجل في هذه المادة", 404)
        if status == "absent":
            db().execute("DELETE FROM attendance WHERE session_id=? AND student_id=?", (sid, d["student_id"]))
        else:
            db().execute(
                "INSERT INTO attendance(session_id,student_id,status,at,method) VALUES(?,?,?,?,'manual') "
                "ON CONFLICT(session_id,student_id) DO UPDATE SET status=excluded.status, method='manual'",
                (sid, d["student_id"], status, now()))
        db().commit()
        return jsonify({"ok": True})

    @app.post("/api/doctor/sessions/<int:sid>/end")
    @auth("doctor", "admin")
    def end_session(sid):
        s = load_session(sid)
        if not s:
            return err("not_found", "المحاضرة غير موجودة", 404)
        if s["ended_at"] is None:
            db().execute("UPDATE sessions SET ended_at=? WHERE id=?", (now(), sid))
            db().commit()
        return jsonify({"ok": True})

    @app.get("/api/doctor/courses/<int:cid>/report")
    @auth("doctor", "admin")
    def report(cid):
        c = db().execute("SELECT * FROM courses WHERE id=?", (cid,)).fetchone()
        if not c or not can_access_course(g.user, cid):
            return err("not_found", "المادة غير موجودة", 404)
        sessions, rows, _ = course_report(c)
        return jsonify({"course": course_dict(c), "sessions": len(sessions), "rows": rows,
                        "unassigned": unassigned(c), "limit_pct": setting("absence_limit_pct")})

    @app.get("/api/doctor/courses/<int:cid>/sessions")
    @auth("doctor", "admin")
    def course_sessions(cid):
        if not can_access_course(g.user, cid):
            return err("forbidden", "غير مسموح", 403)
        out = []
        for s in db().execute("SELECT * FROM sessions WHERE course_id=? ORDER BY started_at DESC LIMIT 60", (cid,)):
            n = db().execute("SELECT COUNT(*) n FROM attendance WHERE session_id=?", (s["id"],)).fetchone()["n"]
            out.append({"id": s["id"], "started_at": s["started_at"], "ended_at": s["ended_at"],
                        "shift": s["shift"], "attended": n})
        return jsonify(out)

    # --------------------------------------------------------------- التصدير
    @app.post("/api/export-link")
    @auth("doctor", "admin")
    def export_link():
        cid = body().get("course_id")
        c = db().execute("SELECT id FROM courses WHERE id=?", (cid,)).fetchone()
        if not c or not can_access_course(g.user, c["id"]):
            return err("forbidden", "غير مسموح", 403)
        t = export_s.dumps({"c": c["id"]})
        return jsonify({"path": f"/export/course/{c['id']}.xlsx?t={t}"})

    @app.get("/export/course/<int:cid>.xlsx")
    def export_course(cid):
        try:
            data = export_s.loads(request.args.get("t", ""), max_age=180)
        except (BadSignature, SignatureExpired):
            return err("expired", "انتهت صلاحية رابط التصدير، أعد الطلب من التطبيق", 403)
        if data.get("c") != cid:
            return err("forbidden", "رابط غير صالح", 403)
        c = db().execute("SELECT * FROM courses WHERE id=?", (cid,)).fetchone()
        if not c:
            return err("not_found", "غير موجود", 404)
        return send_file(build_xlsx(c), as_attachment=True,
                         download_name=f"attendance_{c['code']}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def build_xlsx(course):
        import datetime as dt
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        sessions, rows, recs = course_report(course)
        wb = Workbook()
        ws = wb.active
        ws.title = "كشف الحضور"
        ws.sheet_view.rightToLeft = True
        ws.append([f"جامعة المعقل الأهلية – البصرة | {course['name_ar']} ({course['code']}) | {course['dept_name']}"])
        ws.append([])
        sym = {"present": "ح", "late": "م", "excused": "ع", "absent": "غ"}
        head = ["م", "الرقم المعرفي", "اسم الطالب", "الدوام"] + [
            dt.datetime.fromtimestamp(s["started_at"] + 3 * 3600, dt.timezone.utc).strftime("%m/%d") for s in sessions
        ] + ["حاضر", "متأخر", "بعذر", "غائب", "النسبة %"]
        ws.append(head)
        for cell in ws[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4338CA")
            cell.alignment = Alignment(horizontal="center")
        ws["A1"].font = Font(bold=True, size=13)
        for i, r in enumerate(rows, 1):
            srec = recs.get(r["student_id"], {})
            marks = []
            for s in sessions:
                if s["shift"] not in ("all", r["shift"]):
                    marks.append("-")
                else:
                    marks.append(sym[srec[s["id"]]["status"]] if s["id"] in srec else sym["absent"])
            ws.append([i, r["uid"], r["name"], "صباحي" if r["shift"] == "morning" else "مسائي"] + marks +
                      [r["present"], r["late"], r["excused"], r["absent"], r["pct"] if r["pct"] is not None else ""])
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 34
        ws.append([])
        ws.append(["الرموز: ح حاضر | م متأخر | ع غياب بعذر | غ غائب"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    # =================================================================== المدير
    @app.get("/api/admin/overview")
    @auth("admin")
    def admin_overview():
        autoclose()
        q = lambda sql, *a: db().execute(sql, a).fetchone()[0]
        today = now() - (now() + 3 * 3600) % 86400  # بداية اليوم بتوقيت بغداد
        return jsonify({
            "students": q("SELECT COUNT(*) FROM users WHERE role='student'"),
            "activated": q("SELECT COUNT(*) FROM users WHERE role='student' AND activated=1"),
            "doctors": q("SELECT COUNT(*) FROM users WHERE role='doctor' AND active=1"),
            "courses": q("SELECT COUNT(*) FROM courses WHERE season=?", setting("current_season")),
            "live_sessions": q("SELECT COUNT(*) FROM sessions WHERE ended_at IS NULL"),
            "sessions_today": q("SELECT COUNT(*) FROM sessions WHERE started_at>=?", today),
            "checkins_today": q("SELECT COUNT(*) FROM attendance WHERE at>=?", today),
            "settings": {k: setting(k) for k in DEFAULT_SETTINGS},
        })

    @app.get("/api/admin/settings")
    @auth("admin")
    def get_settings():
        return jsonify({k: setting(k) for k in DEFAULT_SETTINGS})

    @app.post("/api/admin/settings")
    @auth("admin")
    def set_settings():
        d = body()
        ranges = {"current_season": (1, 2), "late_minutes": (1, 120), "absence_limit_pct": (1, 100), "qr_step": (10, 60)}
        for k, v in d.items():
            if k not in ranges:
                continue
            try:
                iv = int(v)
            except (TypeError, ValueError):
                return err("bad_value", "قيمة غير صحيحة")
            lo, hi = ranges[k]
            if not lo <= iv <= hi:
                return err("bad_value", f"القيمة يجب أن تكون بين {lo} و{hi}")
            db().execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                         (k, str(iv)))
        db().commit()
        return jsonify({k: setting(k) for k in DEFAULT_SETTINGS})

    @app.get("/api/admin/doctors")
    @auth("admin")
    def list_doctors():
        out = []
        for u in db().execute("SELECT * FROM users WHERE role='doctor' ORDER BY name"):
            n = db().execute("SELECT COUNT(*) n FROM assignments WHERE doctor_id=?", (u["id"],)).fetchone()["n"]
            out.append({"id": u["id"], "name": u["name"], "email": u["email"], "active": bool(u["active"]), "courses": n})
        return jsonify(out)

    @app.post("/api/admin/doctors")
    @auth("admin")
    def add_doctor():
        d = body()
        name = re.sub(r"\s+", " ", (d.get("name") or "")).strip()
        email = norm_email(d.get("email"))
        pw = d.get("password") or ""
        if len(name) < 3:
            return err("bad_name", "اكتب اسم الدكتور")
        if not email:
            return err("bad_email", "صيغة البريد الإلكتروني غير صحيحة")
        if len(pw) < 6:
            return err("weak_password", "كلمة المرور المؤقتة يجب ألا تقل عن 6 أحرف")
        if db().execute("SELECT 1 FROM users WHERE email=?", (email,)).fetchone():
            return err("email_taken", "هذا البريد مستخدم لحساب آخر", 409)
        cur = db().execute(
            "INSERT INTO users(role,email,name,pw_hash,activated,must_change,created_at) VALUES('doctor',?,?,?,1,1,?)",
            (email, name, generate_password_hash(pw), now()))
        db().commit()
        return jsonify({"id": cur.lastrowid})

    @app.post("/api/admin/doctors/<int:did>/reset-password")
    @auth("admin")
    def doctor_reset(did):
        pw = body().get("password") or ""
        if len(pw) < 6:
            return err("weak_password", "كلمة المرور المؤقتة يجب ألا تقل عن 6 أحرف")
        r = db().execute(
            "UPDATE users SET pw_hash=?, must_change=1, token_ver=token_ver+1 WHERE id=? AND role='doctor'",
            (generate_password_hash(pw), did))
        db().commit()
        return jsonify({"ok": r.rowcount == 1})

    @app.post("/api/admin/doctors/<int:did>/active")
    @auth("admin")
    def doctor_active(did):
        on = 1 if body().get("active") else 0
        db().execute("UPDATE users SET active=?, token_ver=token_ver+1 WHERE id=? AND role='doctor'", (on, did))
        db().commit()
        return jsonify({"ok": True})

    @app.get("/api/admin/courses")
    @auth("admin")
    def admin_courses():
        did = request.args.get("doctor_id", type=int)
        season = request.args.get("season", default=setting("current_season"), type=int)
        rows = db().execute("SELECT * FROM courses WHERE season=? ORDER BY dept_name, stage, name_ar", (season,)).fetchall()
        assigned = set()
        if did:
            assigned = {r["course_id"] for r in db().execute("SELECT course_id FROM assignments WHERE doctor_id=?", (did,))}
        return jsonify([{**course_dict(c), "assigned": c["id"] in assigned} for c in rows])

    @app.post("/api/admin/assign")
    @auth("admin")
    def assign():
        d = body()
        if not db().execute("SELECT 1 FROM users WHERE id=? AND role='doctor'", (d.get("doctor_id"),)).fetchone():
            return err("not_found", "الدكتور غير موجود", 404)
        if not db().execute("SELECT 1 FROM courses WHERE id=?", (d.get("course_id"),)).fetchone():
            return err("not_found", "المادة غير موجودة", 404)
        if d.get("assigned"):
            db().execute("INSERT OR IGNORE INTO assignments(doctor_id,course_id) VALUES(?,?)", (d["doctor_id"], d["course_id"]))
        else:
            db().execute("DELETE FROM assignments WHERE doctor_id=? AND course_id=?", (d["doctor_id"], d["course_id"]))
        db().commit()
        return jsonify({"ok": True})

    @app.get("/api/admin/students")
    @auth("admin")
    def admin_students():
        qtxt = (request.args.get("q") or "").strip()
        sql, args = "SELECT * FROM users WHERE role='student'", []
        if qtxt:
            sql += " AND (uid LIKE ? OR name LIKE ?)"
            args += [f"%{qtxt}%", f"%{qtxt}%"]
        for col in ("stage", "shift", "department"):
            v = request.args.get(col)
            if v:
                sql += f" AND {col}=?"
                args.append(v)
        total = db().execute(sql.replace("SELECT *", "SELECT COUNT(*)", 1), args).fetchone()[0]
        rows = db().execute(sql + " ORDER BY name LIMIT 40", args).fetchall()
        return jsonify({"total": total, "rows": [{
            "id": r["id"], "uid": r["uid"], "name": r["name"], "stage": r["stage"], "shift": r["shift"],
            "department": r["department"], "email": r["email"], "activated": bool(r["activated"]),
            "has_device": bool(r["device_id"]), "active": bool(r["active"])} for r in rows]})

    @app.post("/api/admin/students/<int:sid>/reset")
    @auth("admin")
    def student_reset(sid):
        what = body().get("what")
        if what == "device":
            db().execute("UPDATE users SET device_id=NULL, token_ver=token_ver+1 WHERE id=? AND role='student'", (sid,))
        elif what == "account":
            db().execute(
                "UPDATE users SET pw_hash=NULL, activated=0, device_id=NULL, email=NULL, token_ver=token_ver+1 "
                "WHERE id=? AND role='student'", (sid,))
        else:
            return err("bad_request", "إجراء غير معروف")
        db().commit()
        return jsonify({"ok": True})

    @app.post("/api/admin/students/<int:sid>/department")
    @auth("admin")
    def student_department(sid):
        dept = body().get("department")
        if not db().execute("SELECT 1 FROM courses WHERE dept=?", (dept,)).fetchone():
            return err("bad_department", "قسم غير صحيح")
        db().execute("UPDATE users SET department=? WHERE id=? AND role='student'", (dept, sid))
        db().commit()
        return jsonify({"ok": True})

    @app.post("/api/admin/students/bulk-department")
    @auth("admin")
    def bulk_department():
        d = body()
        dept = d.get("department")
        if not db().execute("SELECT 1 FROM courses WHERE dept=?", (dept,)).fetchone():
            return err("bad_department", "قسم غير صحيح")
        sql, args = "UPDATE users SET department=? WHERE role='student'", [dept]
        if d.get("stage"):
            sql += " AND stage=?"
            args.append(d["stage"])
        if d.get("shift"):
            sql += " AND shift=?"
            args.append(d["shift"])
        if not d.get("overwrite"):
            sql += " AND (department IS NULL OR department='')"
        r = db().execute(sql, args)
        db().commit()
        return jsonify({"updated": r.rowcount})

    return app


# ---------------------------------------------------------------- تهيئة
def _load_secret(folder):
    env = os.environ.get("MAQAL_SECRET")
    if env:
        return env
    path = os.path.join(folder, "secret.key")
    if not os.path.exists(path):
        fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, "w") as f:
            f.write(secrets.token_hex(32))
    with open(path) as f:
        return f.read().strip()


def init_db(path):
    conn = sqlite3.connect(path)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.executescript(SCHEMA)
    for k, v in DEFAULT_SETTINGS.items():
        conn.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", (k, v))
    conn.commit()
    conn.close()


def import_courses(path_db, courses):
    conn = sqlite3.connect(path_db)
    for c in courses:
        conn.execute(
            "INSERT INTO courses(ref_id,dept,dept_name,stage,semester,season,code,name_ar,name_en,ects) "
            "VALUES(:ref_id,:dept,:dept_name,:stage,:semester,:season,:code,:name_ar,:name_en,:ects) "
            "ON CONFLICT(ref_id) DO UPDATE SET dept_name=excluded.dept_name, name_ar=excluded.name_ar, "
            "name_en=excluded.name_en, ects=excluded.ects, semester=excluded.semester, season=excluded.season", c)
    conn.commit()
    conn.close()


def import_students(path_db, students):
    conn = sqlite3.connect(path_db)
    for s in students:
        conn.execute(
            "INSERT INTO users(role,uid,name,stage,shift,created_at) VALUES('student',:uid,:name,:stage,:shift,:ts) "
            "ON CONFLICT(uid) DO UPDATE SET name=excluded.name, stage=excluded.stage, shift=excluded.shift",
            {**s, "ts": now()})
    conn.commit()
    conn.close()


def import_departments(path_db, pairs):
    """pairs: [(uid, dept_key)] — لتعيين قسم كل طالب دفعة واحدة عند توفر البيانات."""
    conn = sqlite3.connect(path_db)
    n = 0
    for uid, dept in pairs:
        n += conn.execute("UPDATE users SET department=? WHERE role='student' AND uid=?", (dept, uid)).rowcount
    conn.commit()
    conn.close()
    return n


def create_admin(path_db, email, name, password):
    conn = sqlite3.connect(path_db)
    conn.execute(
        "INSERT INTO users(role,email,name,pw_hash,activated,created_at) VALUES('admin',?,?,?,1,?) "
        "ON CONFLICT(email) DO UPDATE SET pw_hash=excluded.pw_hash, token_ver=token_ver+1",
        (email.lower(), name, generate_password_hash(password), now()))
    conn.commit()
    conn.close()


if __name__ == "__main__":
    create_app().run(host="0.0.0.0", port=int(os.environ.get("PORT", 8000)))
